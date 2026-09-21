"""
Excel Export Engine for AIET Employee Attendance & Salary Processing System.
Utilizes openpyxl to generate beautifully styled spreadsheets with filters, freeze panes,
currency formatting, and auto-adjusted column widths.
"""

import os
import csv
from datetime import datetime
import models
from utils import get_reports_dir, get_month_name

def export_employees_to_excel(output_path=None):
    """
    Export full employee directory into a formatted Excel sheet.
    """
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.join(get_reports_dir(), f"AIET_Employees_Export_{timestamp}.xlsx")

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        # Fallback to CSV if openpyxl not installed
        csv_path = output_path.replace(".xlsx", ".csv")
        return export_employees_to_csv(csv_path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees Directory"
    ws.views.sheetView[0].showGridLines = True

    # Institutional Header
    header_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    sub_font = Font(name="Calibri", size=10, italic=True, color="475569")
    ws.merge_cells("A1:L1")
    ws["A1"] = "ALVAS INSTITUTE OF ENGINEERING AND TECHNOLOGY (AIET)"
    ws["A1"].font = header_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:L2")
    ws["A2"] = f"Master Employee Register — Exported On: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = sub_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # Column Headers
    headers = [
        "Employee ID", "Full Name", "Department", "Designation", "Phone Number",
        "Email Address", "Joining Date", "Basic Salary (₹)", "OT Rate (₹/hr)",
        "Allowances (₹)", "Deductions (₹)", "Status"
    ]

    col_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    col_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    header_row = 4
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.value = header_title
        cell.font = col_header_font
        cell.fill = col_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Data Rows
    employees = models.get_all_employees()
    currency_format = '₹ #,##0.00'

    current_row = header_row + 1
    for emp in employees:
        row_values = [
            emp["employee_id"],
            emp["name"],
            emp["department"],
            emp["designation"],
            emp["phone"],
            emp["email"],
            emp["joining_date"],
            emp["basic_salary"],
            emp["overtime_rate"],
            emp["allowances"],
            emp["deductions"],
            emp["status"]
        ]

        for col_num, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.value = val
            cell.border = thin_border

            # Number formatting
            if col_num in (8, 9, 10, 11):
                cell.number_format = currency_format
                cell.alignment = Alignment(horizontal="right")
            elif col_num in (1, 5, 7, 12):
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")

        current_row += 1

    # Freeze panes below header
    ws.freeze_panes = ws[f"A{header_row + 1}"]

    # Enable filters
    last_col_letter = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{current_row - 1}"

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.row < header_row:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)
    return True, output_path

def export_attendance_to_excel(month=None, year=None, output_path=None):
    """
    Export attendance records to formatted Excel workbook.
    """
    if not output_path:
        m_str = f"Month_{month}" if month else "AllMonths"
        y_str = str(year) if year else "AllYears"
        output_path = os.path.join(get_reports_dir(), f"AIET_Attendance_{m_str}_{y_str}.xlsx")

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        csv_path = output_path.replace(".xlsx", ".csv")
        return export_attendance_to_csv(month, year, csv_path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Monthly Attendance"
    ws.views.sheetView[0].showGridLines = True

    period_str = f"{get_month_name(month)} {year}" if month and year else "All Recorded Periods"

    ws.merge_cells("A1:I1")
    ws["A1"] = "ALVAS INSTITUTE OF ENGINEERING AND TECHNOLOGY (AIET)"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:I2")
    ws["A2"] = f"Monthly Attendance Register — Period: {period_str}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="475569")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        "Att. ID", "Employee ID", "Employee Name", "Department",
        "Month / Year", "Total Working Days", "Days Present", "Leave Days", "OT Hours"
    ]

    col_header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    col_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    header_row = 4
    for col_num, h in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col_num, value=h)
        c.font = col_header_font
        c.fill = col_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    records = models.get_all_attendance(month, year)
    current_row = header_row + 1

    for r in records:
        m_name = get_month_name(r["month"])
        row_values = [
            r["attendance_id"],
            r["employee_id"],
            r["name"],
            r["department"],
            f"{m_name} {r['year']}",
            r["total_working_days"],
            r["days_present"],
            r["leave_days"],
            r["overtime_hours"]
        ]

        for col_num, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_num, value=val)
            cell.border = thin_border
            if col_num in (1, 5, 6, 7, 8, 9):
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")
        current_row += 1

    ws.freeze_panes = ws[f"A{header_row + 1}"]
    last_col_letter = get_column_letter(len(headers))
    if current_row > header_row + 1:
        ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{current_row - 1}"

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.row < header_row:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)
    return True, output_path

def export_salary_report_to_excel(month=None, year=None, output_path=None):
    """
    Export processed monthly payroll and salary records into a master Excel report.
    """
    if not output_path:
        m_str = f"Month_{month}" if month else "AllMonths"
        y_str = str(year) if year else "AllYears"
        output_path = os.path.join(get_reports_dir(), f"AIET_Salary_Payroll_{m_str}_{y_str}.xlsx")

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        csv_path = output_path.replace(".xlsx", ".csv")
        return export_salary_to_csv(month, year, csv_path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Salary Payroll Report"
    ws.views.sheetView[0].showGridLines = True

    period_str = f"{get_month_name(month)} {year}" if month and year else "All Processed Periods"

    ws.merge_cells("A1:K1")
    ws["A1"] = "ALVAS INSTITUTE OF ENGINEERING AND TECHNOLOGY (AIET)"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:K2")
    ws["A2"] = f"Consolidated Staff Salary & Payroll Report — Period: {period_str}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="475569")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        "Employee ID", "Employee Name", "Department", "Period",
        "Basic Salary (₹)", "Days Present", "Attendance Salary (₹)",
        "OT Pay (₹)", "Allowances (₹)", "Gross Salary (₹)", "Deductions (₹)", "Net Payable (₹)"
    ]

    col_header_fill = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
    col_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    header_row = 4
    for col_num, h in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col_num, value=h)
        c.font = col_header_font
        c.fill = col_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    records = models.get_all_salary_records(month, year)
    currency_format = '₹ #,##0.00'
    current_row = header_row + 1

    tot_gross = 0.0
    tot_deduct = 0.0
    tot_net = 0.0

    for r in records:
        m_name = get_month_name(r["month"])
        row_values = [
            r["employee_id"],
            r["name"],
            r["department"],
            f"{m_name} {r['year']}",
            r["basic_salary"],
            f"{r['days_present']} / {r['total_working_days']}",
            r["attendance_salary"],
            r["overtime_pay"],
            r["allowances"],
            r["gross_salary"],
            r["deductions"],
            r["net_salary"]
        ]

        tot_gross += float(r["gross_salary"])
        tot_deduct += float(r["deductions"])
        tot_net += float(r["net_salary"])

        for col_num, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_num, value=val)
            cell.border = thin_border
            if col_num in (5, 7, 8, 9, 10, 11, 12):
                cell.number_format = currency_format
                cell.alignment = Alignment(horizontal="right")
            elif col_num in (1, 4, 6):
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")
        current_row += 1

    # Totals Summary Row
    if records:
        total_row = current_row
        ws.cell(row=total_row, column=1, value="TOTALS").font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=total_row, column=10, value=tot_gross).number_format = currency_format
        ws.cell(row=total_row, column=10).font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=total_row, column=11, value=tot_deduct).number_format = currency_format
        ws.cell(row=total_row, column=11).font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=total_row, column=12, value=tot_net).number_format = currency_format
        ws.cell(row=total_row, column=12).font = Font(name="Calibri", size=11, bold=True, color="15803D")

        for col_num in range(1, len(headers) + 1):
            ws.cell(row=total_row, column=col_num).fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            ws.cell(row=total_row, column=col_num).border = thin_border
        current_row += 1

    ws.freeze_panes = ws[f"A{header_row + 1}"]
    last_col_letter = get_column_letter(len(headers))
    if len(records) > 0:
        ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{current_row - 2}"

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.row < header_row:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)
    return True, output_path

# ==============================================================================
# CSV Fallback Handlers
# ==============================================================================

def export_employees_to_csv(csv_path):
    employees = models.get_all_employees()
    with open(csv_path, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(["Employee ID", "Name", "Department", "Designation", "Phone", "Email", "Joining Date", "Basic Salary", "Overtime Rate", "Allowances", "Deductions", "Status"])
        for emp in employees:
            writer.writerow([
                emp["employee_id"], emp["name"], emp["department"], emp["designation"],
                emp["phone"], emp["email"], emp["joining_date"], emp["basic_salary"],
                emp["overtime_rate"], emp["allowances"], emp["deductions"], emp["status"]
            ])
    return True, csv_path

def export_attendance_to_csv(month, year, csv_path):
    records = models.get_all_attendance(month, year)
    with open(csv_path, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(["Att ID", "Emp ID", "Name", "Department", "Month", "Year", "Working Days", "Days Present", "Leave Days", "OT Hours"])
        for r in records:
            writer.writerow([
                r["attendance_id"], r["employee_id"], r["name"], r["department"],
                r["month"], r["year"], r["total_working_days"], r["days_present"],
                r["leave_days"], r["overtime_hours"]
            ])
    return True, csv_path

def export_salary_to_csv(month, year, csv_path):
    records = models.get_all_salary_records(month, year)
    with open(csv_path, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(["Emp ID", "Name", "Department", "Month", "Year", "Basic Salary", "Working Days", "Days Present", "Att Salary", "OT Pay", "Allowances", "Gross Salary", "Deductions", "Net Salary"])
        for r in records:
            writer.writerow([
                r["employee_id"], r["name"], r["department"], r["month"], r["year"],
                r["basic_salary"], r["total_working_days"], r["days_present"],
                r["attendance_salary"], r["overtime_pay"], r["allowances"],
                r["gross_salary"], r["deductions"], r["net_salary"]
            ])
    return True, csv_path
